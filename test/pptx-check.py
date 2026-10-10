"""Confere os .pptx gerados em test/out/pptx-*.pptx: abre com python-pptx, conta slides, gráficos nativos, planilhas embutidas e textos.
Uso: python3 test/pptx-check.py   (gere antes com test/m/pptx-make.js, hash 'arquivo|aba|tipo1,tipo2')"""
import glob, io, sys, zipfile
from pptx import Presentation
bad = []
for f in sorted(glob.glob('test/out/pptx-*.pptx')):
    name = f.split('pptx-')[1][:-5]
    try:
        p = Presentation(f)
        charts = texts = tables = 0; wb = 0; groups = 0
        for s in p.slides:
            for sh in s.shapes:
                if sh.has_chart:
                    charts += 1; ch = sh.chart
                    xl = ch.part.chart_workbook.xlsx_part
                    if xl is not None and len(xl.blob) > 200:
                        wb += 1
                        z = zipfile.ZipFile(io.BytesIO(xl.blob)); assert 'xl/worksheets/sheet1.xml' in z.namelist()
                    for pl in ch.plots:
                        for se in pl.series:
                            vals = list(se.values) if hasattr(se, 'values') else []
                            assert len(vals) > 0 or name in ('kpi',), 'série vazia'
                elif sh.shape_type == 6 and len(sh.shapes) > 20: groups += 1  # desenho vetorial agrupado (árvore radial)
                elif sh.has_text_frame and sh.text_frame.text.strip(): texts += 1
                elif getattr(sh, 'has_table', False) and sh.has_table: tables += 1
        ok = len(p.slides) >= 4 and texts >= 8 and (charts == wb) and (charts >= 1 or groups >= 1 or name == 'kpi')
        print(('OK   ' if ok else 'FALHA'), name.ljust(11), 'slides', len(p.slides), 'gráficos', charts, 'planilhas', wb, 'textos', texts, 'tabelas', tables, 'desenhos', groups)
        if not ok: bad.append(name)
    except Exception as e:
        print('FALHA', name, repr(e)); bad.append(name)
print('FALHAS:', bad if bad else 'nenhuma'); sys.exit(1 if bad else 0)
